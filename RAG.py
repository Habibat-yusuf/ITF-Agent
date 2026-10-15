
import os
from pathlib import Path

import pdfplumber
from docx import Document
from dotenv import load_dotenv
from google import genai
from google.genai import types
from supabase import create_client, Client

# Load environment variables
BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_KEY")

if not GEMINI_API_KEY:
    raise ValueError("GEMINI_API_KEY is missing from .env")

if not SUPABASE_URL or not SUPABASE_KEY:
    raise ValueError("SUPABASE_URL or SUPABASE_SERVICE_KEY is missing")

# Initialize shared clients
client = genai.Client(api_key=GEMINI_API_KEY)
supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

EMBEDDING_MODEL = "gemini-embedding-2-preview"
EMBEDDING_DIMENSIONS = 768
REFUSAL_MESSAGE = (
    "I don't have that information. Please contact ITF directly."
)


def extract_text_from_file(file_path: str) -> str:
    """Extract text from PDF, DOCX, or TXT files."""
    path = Path(file_path)
    ext = path.suffix.lower()
    extracted_parts = []

    if ext == ".pdf":
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    extracted_parts.append(page_text)

    elif ext == ".docx":
        document = Document(path)

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                extracted_parts.append(paragraph.text)

        # Also capture text inside Word tables.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if cells:
                    extracted_parts.append(" | ".join(cells))

    elif ext == ".txt":
        extracted_parts.append(
            path.read_text(encoding="utf-8")
        )

    else:
        raise ValueError(
            f"Unsupported file format: {ext}. "
            "Upload a PDF, DOCX, or TXT file."
        )

    return "\n".join(extracted_parts).strip()


def save_extracted_text(
    text: str,
    source_name: str,
    output_dir: str = "extracted_text",
) -> str:
    """Save extracted text to a UTF-8 text file."""
    output_path = BASE_DIR / output_dir
    output_path.mkdir(parents=True, exist_ok=True)

    safe_name = Path(source_name).stem
    text_file = output_path / f"{safe_name}.txt"
    text_file.write_text(text, encoding="utf-8")

    return str(text_file)


def chunk_text(
    text: str,
    chunk_size: int = 500,
    overlap: int = 50,
) -> list[str]:
    """Split text into overlapping character chunks."""
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError(
            "overlap must be non-negative and smaller than chunk_size"
        )

    chunks = []
    step = chunk_size - overlap

    for start in range(0, len(text), step):
        chunk = text[start:start + chunk_size].strip()
        if chunk:
            chunks.append(chunk)

    return chunks


def get_embedding(text: str) -> list[float]:
    """Generate an embedding for text."""
    response = client.models.embed_content(
        model=EMBEDDING_MODEL,
        contents=text.replace("\n", " "),
        config=types.EmbedContentConfig(
            output_dimensionality=EMBEDDING_DIMENSIONS
        ),
    )

    if not response.embeddings:
        raise RuntimeError("The embedding API returned no embedding.")

    return response.embeddings[0].values


def process_and_store_document(
    file_path: str,
    source_name: str | None = None,
) -> dict:
    """Extract, save, chunk, embed, and store a document."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    print(f"Processing {path.name}...")

    raw_text = extract_text_from_file(str(path))

    if not raw_text:
        raise ValueError(
            "No text could be extracted. The document may be scanned "
            "or contain no readable text."
        )

    saved_text_path = save_extracted_text(
        raw_text,
        source_name or path.name,
    )

    chunks = chunk_text(raw_text)
    records = []

    for chunk in chunks:
        clean_chunk = chunk.replace("\x00", "").strip()
        if not clean_chunk:
            continue

        embedding = get_embedding(clean_chunk)

        records.append({
            "content": clean_chunk,
            "source": source_name or path.name,
            "embedding": embedding,
        })

    if not records:
        raise ValueError("No usable text chunks were generated.")

    # Insert in batches to reduce the number of database requests.
    batch_size = 50

    for start in range(0, len(records), batch_size):
        batch = records[start:start + batch_size]
        supabase.table("itf_documents").insert(batch).execute()

    return {
        "source": source_name or path.name,
        "chunks_stored": len(records),
        "text_file": saved_text_path,
    }


def retrieve_context(
    query: str,
    match_count: int = 3,
    match_threshold: float = 0.1,
) -> dict:
    """Retrieve relevant document chunks from Supabase."""
    if not query.strip():
        return {"context": "", "sources": [], "chunks": []}

    query_embedding = get_embedding(query)

    response = supabase.rpc(
        "match_itf_documents",
        {
            "query_embedding": query_embedding,
            "match_threshold": match_threshold,
            "match_count": match_count,
        },
    ).execute()

    chunks = response.data or []

    context = "\n\n".join(
        row["content"]
        for row in chunks
        if row.get("content")
    )

    sources = sorted({
        row["source"]
        for row in chunks
        if row.get("source")
    })

    return {
        "context": context,
        "sources": sources,
        "chunks": chunks,
    }


if __name__ == "__main__":
    brochure_path = BASE_DIR / "ITF BROCHURE.pdf"

    result = process_and_store_document(
        str(brochure_path),
        "ITF 2022 Brochure",
    )

    print("Ingestion completed:", result)
